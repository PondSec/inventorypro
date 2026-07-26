import json
import io
import tempfile
import zipfile
from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook

import app as inventory_app


class ServerSettingsTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        temp_path = Path(self.temp_dir.name)
        inventory_app.DATABASE = str(temp_path / "test_inventory.db")
        inventory_app.UPLOADS_DIR = temp_path / "uploads"
        inventory_app.APP_INSTANCE_PATH = temp_path / "instance"
        inventory_app.RUNTIME_CONFIG_PATH = temp_path / "runtime_config.json"
        inventory_app.RUNTIME_SETTINGS_CACHE = {
            "host": "0.0.0.0",
            "port": 5000,
            "debug": False
        }
        with inventory_app.app.app_context():
            inventory_app.init_db()
            db = inventory_app.get_db()
            password_hash = inventory_app.generate_password_hash("secret1234")
            cursor = db.execute(
                "INSERT INTO users (username, password_hash) VALUES (?, ?)",
                ("tester", password_hash)
            )
            admin_role = db.execute("SELECT id FROM roles WHERE name = 'Admin'").fetchone()
            if admin_role:
                inventory_app.assign_user_role(db, cursor.lastrowid, "Admin")
            category_id = db.execute(
                "INSERT INTO categories (name) VALUES (?)",
                ("Testkategorie",),
            ).lastrowid
            self.device_id = db.execute(
                "INSERT INTO devices (name, category_id, serial_number, specs) VALUES (?, ?, ?, ?)",
                ("Testgerät", category_id, "SER-001", "{}"),
            ).lastrowid
            db.commit()
        self.client = inventory_app.app.test_client()

    def tearDown(self):
        self.temp_dir.cleanup()

    def login(self, username="tester", password="secret1234"):
        return self.client.post("/login", data={"username": username, "password": password})

    def test_get_and_update_settings(self):
        self.login()
        response = self.client.get("/api/settings/server")
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        settings = payload["settings"]
        settings["server"]["host"] = "127.0.0.1"
        update_response = self.client.put("/api/settings/server", json=settings)
        self.assertEqual(update_response.status_code, 200)
        updated = update_response.get_json()
        self.assertTrue(updated["meta"]["pendingRestart"])

    def test_settings_page_links_to_personal_account_security(self):
        self.login()

        response = self.client.get("/settings")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'href="/account/security"', response.data)
        self.assertIn(b"Erfordert TOTP f\xc3\xbcr alle Konten beim n\xc3\xa4chsten Login.", response.data)

    def test_invalid_port_rejected(self):
        self.login()
        response = self.client.get("/api/settings/server")
        settings = response.get_json()["settings"]
        settings["server"]["port"] = 70000
        update_response = self.client.put("/api/settings/server", json=settings)
        self.assertEqual(update_response.status_code, 400)

    def test_signed_update_policy_is_validated_and_published(self):
        self.login()
        response = self.client.get("/api/settings/server")
        settings = response.get_json()["settings"]
        settings["updates"] = {
            "autoUpdateEnabled": True,
            "channel": "stable",
            "checkIntervalMinutes": 120,
            "maintenanceWindow": "02:45",
        }

        update_response = self.client.put("/api/settings/server", json=settings)

        self.assertEqual(update_response.status_code, 200)
        policy_path = inventory_app.APP_INSTANCE_PATH / "update_policy.json"
        self.assertTrue(policy_path.exists())
        policy = json.loads(policy_path.read_text(encoding="utf-8"))
        self.assertTrue(policy["autoUpdateEnabled"])
        self.assertEqual(policy["channel"], "stable")
        self.assertEqual(policy["checkIntervalMinutes"], 120)
        self.assertEqual(policy["maintenanceWindow"], "02:45")
        self.assertFalse(any("password" in key.lower() for key in policy))

    def test_update_maintenance_window_is_normalized_before_publishing(self):
        self.login()
        for raw_window, normalized_window in (
            ("09:00", "09:00"),
            ("9:00", "09:00"),
            ("09:00:00", "09:00"),
            ("23:30", "23:30"),
            ("00:00", "00:00"),
        ):
            with self.subTest(raw_window=raw_window):
                settings = self.client.get("/api/settings/server").get_json()["settings"]
                settings["updates"] = {
                    "autoUpdateEnabled": True,
                    "channel": "stable",
                    "checkIntervalMinutes": 120,
                    "maintenanceWindow": raw_window,
                }

                response = self.client.put("/api/settings/server", json=settings)

                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    response.get_json()["settings"]["updates"]["maintenanceWindow"],
                    normalized_window,
                )
                policy = json.loads(
                    (inventory_app.APP_INSTANCE_PATH / "update_policy.json").read_text(encoding="utf-8")
                )
                self.assertEqual(policy["maintenanceWindow"], normalized_window)

    def test_update_maintenance_window_with_seconds_is_rejected(self):
        self.login()
        settings = self.client.get("/api/settings/server").get_json()["settings"]
        settings["updates"] = {
            "autoUpdateEnabled": True,
            "channel": "stable",
            "checkIntervalMinutes": 120,
            "maintenanceWindow": "09:00:30",
        }

        response = self.client.put("/api/settings/server", json=settings)

        self.assertEqual(response.status_code, 400)
        self.assertIn("updates.maintenanceWindow", response.get_json()["details"])

    def test_disabled_updates_do_not_block_server_settings_save(self):
        self.login()
        response = self.client.get("/api/settings/server")
        settings = response.get_json()["settings"]
        settings["updates"] = {
            "autoUpdateEnabled": False,
            "channel": "preview",
            "checkIntervalMinutes": "invalid",
            "maintenanceWindow": "",
        }

        update_response = self.client.put("/api/settings/server", json=settings)

        self.assertEqual(update_response.status_code, 200)
        policy_path = inventory_app.APP_INSTANCE_PATH / "update_policy.json"
        policy = json.loads(policy_path.read_text(encoding="utf-8"))
        self.assertFalse(policy["autoUpdateEnabled"])
        self.assertEqual(policy["channel"], "stable")
        self.assertEqual(policy["checkIntervalMinutes"], 360)
        self.assertEqual(policy["maintenanceWindow"], "03:30")

    def test_unsigned_update_channel_is_rejected(self):
        self.login()
        response = self.client.get("/api/settings/server")
        settings = response.get_json()["settings"]
        settings["updates"]["autoUpdateEnabled"] = True
        settings["updates"]["channel"] = "preview"

        update_response = self.client.put("/api/settings/server", json=settings)

        self.assertEqual(update_response.status_code, 400)
        self.assertIn("updates.channel", update_response.get_json()["details"])

    def test_backup_run(self):
        self.login()
        response = self.client.get("/api/settings/server")
        settings = response.get_json()["settings"]
        settings["backup"]["enabled"] = True
        settings["backup"]["directory"] = str(Path(self.temp_dir.name) / "backups")
        update_response = self.client.put("/api/settings/server", json=settings)
        self.assertEqual(update_response.status_code, 200)
        run_response = self.client.post("/api/backups/run", json={"force": True})
        self.assertEqual(run_response.status_code, 200)
        result = run_response.get_json()
        self.assertIn(result["status"], {"success", "failed"})

    def test_backup_history_keeps_the_existing_response_shape(self):
        self.login()
        with inventory_app.app.app_context():
            db = inventory_app.get_db()
            db.execute(
                """
                INSERT INTO backup_runs (status, backup_path, backup_size_bytes, message)
                VALUES (?, ?, ?, ?)
                """,
                ("success", "/backups/example.db", 123, "completed"),
            )
            db.commit()

        response = self.client.get("/api/backups/list")
        self.assertEqual(response.status_code, 200)
        backup = response.get_json()["backups"][0]
        self.assertEqual(
            set(backup),
            {"id", "status", "path", "sizeBytes", "message", "createdAt"},
        )
        self.assertEqual(backup["path"], "/backups/example.db")
        self.assertEqual(backup["sizeBytes"], 123)

    def test_tabular_import_preview_and_device_import(self):
        self.login()
        settings = self.client.get("/api/settings/server").get_json()["settings"]
        settings["importExport"]["importAllowed"] = True
        self.assertEqual(self.client.put("/api/settings/server", json=settings).status_code, 200)
        csv_content = b"Hostname;Seriennummer;Standort;Hersteller\ncore-1;SER-1;Berlin;Pond\n"
        preview = self.client.post(
            "/api/import/preview",
            data={"entity": "devices", "file": (io.BytesIO(csv_content), "devices.csv")},
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.get_json()["validRows"], 1)
        preview_data = preview.get_json()
        direct_import = self.client.post(
            "/api/import",
            data={
                "entity": "devices",
                "mode": "append",
                "file": (io.BytesIO(csv_content), "devices.csv"),
            },
        )
        self.assertEqual(direct_import.status_code, 400)
        self.assertIn("Vorschau", direct_import.get_json()["error"])
        imported = self.client.post(
            "/api/import",
            data={
                "entity": "devices",
                "mode": "append",
                "previewToken": preview_data["previewToken"],
                "mapping": json.dumps(preview_data["mapping"]),
                "matchingKey": preview_data["matchingKey"],
                "file": (io.BytesIO(csv_content), "devices.csv"),
            },
        )
        self.assertEqual(imported.status_code, 200)
        self.assertEqual(imported.get_json()["summary"]["created"], 1)

    def test_tabular_preview_includes_complete_downloadable_error_report(self):
        self.login()
        settings = self.client.get("/api/settings/server").get_json()["settings"]
        settings["importExport"]["importAllowed"] = True
        self.assertEqual(self.client.put("/api/settings/server", json=settings).status_code, 200)
        source = b"name,serial\n,missing-1\n,missing-2\nvalid,S-1\n"

        preview = self.client.post(
            "/api/import/preview",
            data={"entity": "devices", "file": (io.BytesIO(source), "devices.csv")},
        )

        self.assertEqual(preview.status_code, 200)
        report = preview.get_json()["errorReport"]
        self.assertEqual(report["format"], "csv")
        self.assertEqual(report["filename"], "inventorypro-devices-import-errors.csv")
        self.assertEqual(report["rowCount"], 2)
        self.assertIn("Zeile,Fehler\n2,Name ist erforderlich.\n", report["content"])

    def test_tabular_preview_reports_duplicates_within_the_uploaded_file(self):
        self.login()
        settings = self.client.get("/api/settings/server").get_json()["settings"]
        settings["importExport"]["importAllowed"] = True
        self.assertEqual(self.client.put("/api/settings/server", json=settings).status_code, 200)
        source = b"name,serial\nfirst,DUP-1\nsecond,DUP-1\n"

        preview = self.client.post(
            "/api/import/preview",
            data={"entity": "devices", "file": (io.BytesIO(source), "devices.csv")},
        )

        self.assertEqual(preview.status_code, 200)
        payload = preview.get_json()
        self.assertEqual(payload["conflictCount"], 1)
        self.assertEqual(
            payload["conflicts"],
            [{"line": 3, "sourceLine": 2, "key": "serial_number", "source": "file"}],
        )

    def test_tabular_xlsx_and_json_previews_are_available(self):
        self.login()
        settings = self.client.get("/api/settings/server").get_json()["settings"]
        settings["importExport"]["importAllowed"] = True
        self.assertEqual(self.client.put("/api/settings/server", json=settings).status_code, 200)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Hostname", "Seriennummer", "Standort"])
        worksheet.append(["xlsx-device", "XLSX-1", "Hamburg"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        xlsx_content = buffer.getvalue()

        preview = self.client.post(
            "/api/import/preview",
            data={"entity": "devices", "file": (io.BytesIO(xlsx_content), "devices.xlsx")},
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.get_json()["format"], "xlsx")
        xlsx_preview = preview.get_json()
        imported = self.client.post(
            "/api/import",
            data={
                "entity": "devices",
                "mode": "append",
                "previewToken": xlsx_preview["previewToken"],
                "mapping": json.dumps(xlsx_preview["mapping"]),
                "matchingKey": xlsx_preview["matchingKey"],
                "sheetName": xlsx_preview["sheetName"],
                "file": (io.BytesIO(xlsx_content), "devices.xlsx"),
            },
        )
        self.assertEqual(imported.status_code, 200)
        self.assertEqual(imported.get_json()["summary"]["created"], 1)

        json_content = b'{"assets": [{"Bezeichnung": "JSON Asset", "Rechnungsnummer": "JSON-1"}]}'
        json_preview = self.client.post(
            "/api/import/preview",
            data={"entity": "assets", "file": (io.BytesIO(json_content), "assets.json")},
        )
        self.assertEqual(json_preview.status_code, 200)
        self.assertEqual(json_preview.get_json()["format"], "json")

    def test_import_profiles_persist_mapping_and_reject_foreign_entities(self):
        self.login()
        settings = self.client.get("/api/settings/server").get_json()["settings"]
        settings["importExport"]["importAllowed"] = True
        self.assertEqual(self.client.put("/api/settings/server", json=settings).status_code, 200)

        invalid_create = self.client.post(
            "/api/import/profiles",
            json={"name": "Ungültig", "entity": "devices", "mapping": {"name": 123}},
        )
        self.assertEqual(invalid_create.status_code, 400)

        created = self.client.post(
            "/api/import/profiles",
            json={
                "name": "Altes Inventar",
                "entity": "devices",
                "mapping": {"name": "Computer", "serial_number": "Asset Tag"},
                "matchingKey": "serial_number",
                "sheetName": "Hardware",
            },
        )
        self.assertEqual(created.status_code, 201)
        profile = created.get_json()["profile"]
        self.assertEqual(profile["mapping"]["name"], "Computer")

        listed = self.client.get("/api/import/profiles?entity=devices")
        self.assertEqual(listed.status_code, 200)
        self.assertEqual([entry["id"] for entry in listed.get_json()["profiles"]], [profile["id"]])
        self.assertEqual(self.client.get("/api/import/profiles?entity=vendors").status_code, 400)

        updated = self.client.put(
            f"/api/import/profiles/{profile['id']}",
            json={
                "name": "Altes Inventar v2",
                "entity": "devices",
                "mapping": {"name": "Computer", "serial_number": "Asset Tag"},
                "matchingKey": "serial_number",
            },
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.get_json()["profile"]["name"], "Altes Inventar v2")
        self.assertEqual(
            self.client.put(
                f"/api/import/profiles/{profile['id']}",
                json={"name": "Ungültig", "entity": "devices", "mapping": {"name": 123}},
            ).status_code,
            400,
        )

        source = b"Computer,Asset Tag\nedge-1,EDGE-1\n"
        preview = self.client.post(
            "/api/import/preview",
            data={
                "entity": "devices",
                "profileId": str(profile["id"]),
                "file": (io.BytesIO(source), "legacy.csv"),
            },
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.get_json()["profileId"], profile["id"])
        self.assertEqual(preview.get_json()["matchingKey"], "serial_number")

        wrong_entity = self.client.post(
            "/api/import/preview",
            data={
                "entity": "assets",
                "profileId": str(profile["id"]),
                "file": (io.BytesIO(source), "legacy.csv"),
            },
        )
        self.assertEqual(wrong_entity.status_code, 400)

        deleted = self.client.delete(f"/api/import/profiles/{profile['id']}")
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(self.client.get("/api/import/profiles").get_json()["profiles"], [])
        self.assertEqual(self.client.delete(f"/api/import/profiles/{profile['id']}").status_code, 404)
        self.assertEqual(self.client.put("/api/import/profiles/999", json={"name": "Fehlt", "entity": "devices", "mapping": {}}).status_code, 404)

    def test_import_profiles_require_administrative_permission(self):
        with inventory_app.app.app_context():
            db = inventory_app.get_db()
            db.execute(
                "INSERT INTO users (username, password_hash) VALUES (?, ?)",
                ("reader", inventory_app.generate_password_hash("reader-password")),
            )
            db.commit()
        self.client.get("/logout")
        self.login("reader", "reader-password")

        response = self.client.get("/api/import/profiles")

        self.assertEqual(response.status_code, 403)

    def test_import_profile_routes_respect_import_disable_switch(self):
        self.login()

        self.assertEqual(self.client.get("/api/import/profiles").status_code, 403)
        self.assertEqual(
            self.client.post("/api/import/profiles", json={"name": "Profil", "entity": "devices", "mapping": {}}).status_code,
            403,
        )
        self.assertEqual(
            self.client.put("/api/import/profiles/1", json={"name": "Profil", "entity": "devices", "mapping": {}}).status_code,
            403,
        )
        self.assertEqual(self.client.delete("/api/import/profiles/1").status_code, 403)

    def test_import_profile_schema_migration_is_recorded_once(self):
        with inventory_app.app.app_context():
            db = inventory_app.get_db()
            migration = db.execute(
                "SELECT id FROM schema_migrations WHERE id = ?",
                ("002_import_profiles",),
            ).fetchall()
            columns = [row["name"] for row in db.execute("PRAGMA table_info(import_profiles)").fetchall()]

        self.assertEqual([row["id"] for row in migration], ["002_import_profiles"])
        self.assertIn("mapping_json", columns)

    def test_xlsx_export_contains_inventory_sheets(self):
        self.login()
        settings = self.client.get("/api/settings/server").get_json()["settings"]
        settings["importExport"]["exportFormat"] = "xlsx"
        settings["importExport"]["includeUploads"] = False
        self.assertEqual(self.client.put("/api/settings/server", json=settings).status_code, 200)

        response = self.client.get("/api/export")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
        self.assertIn("devices", workbook.sheetnames)
        sheet = workbook["devices"]
        self.assertEqual(next(sheet.values)[0], "id")
        self.assertIn("Testgerät", [row[1] for row in sheet.iter_rows(values_only=True)])
        workbook.close()

    def test_inventory_exports_escape_spreadsheet_formulas_and_record_metadata(self):
        self.login()
        with inventory_app.app.app_context():
            db = inventory_app.get_db()
            db.execute("UPDATE devices SET name = ? WHERE id = ?", ("=HYPERLINK(\"https://example.test\")", self.device_id))
            db.commit()

        settings = self.client.get("/api/settings/server").get_json()["settings"]
        settings["importExport"].update({"exportFormat": "csv", "includeUploads": False})
        self.assertEqual(self.client.put("/api/settings/server", json=settings).status_code, 200)
        csv_export = self.client.get("/api/export")
        self.assertEqual(csv_export.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(csv_export.data)) as archive:
            device_csv = archive.read("devices.csv").decode("utf-8")
        self.assertIn("'=HYPERLINK", device_csv)

        settings["importExport"]["exportFormat"] = "xlsx"
        self.assertEqual(self.client.put("/api/settings/server", json=settings).status_code, 200)
        xlsx_export = self.client.get("/api/export")
        workbook = load_workbook(io.BytesIO(xlsx_export.data), read_only=True, data_only=False)
        self.assertIn("metadata", workbook.sheetnames)
        self.assertTrue(
            any(
                isinstance(value, str) and value.startswith("'=HYPERLINK")
                for value in [row[1] for row in workbook["devices"].iter_rows(values_only=True)]
            )
        )
        self.assertEqual(next(workbook["metadata"].iter_rows(values_only=True)), ("key", "value"))
        workbook.close()

        settings["importExport"]["exportFormat"] = "json"
        self.assertEqual(self.client.put("/api/settings/server", json=settings).status_code, 200)
        json_export = self.client.get("/api/export")
        payload = json.loads(json_export.data)
        self.assertEqual(payload["_metadata"]["format"], "json")
        self.assertIn("exportedAt", payload["_metadata"])

        device_export = self.client.get("/api/export/devices")
        self.assertEqual(device_export.status_code, 200)
        self.assertIn("'=HYPERLINK", device_export.get_data(as_text=True))
        with inventory_app.app.app_context():
            db = inventory_app.get_db()
            actions = [row["action"] for row in db.execute("SELECT action FROM activity_log ORDER BY id").fetchall()]
        self.assertIn("export_created", actions)

    def test_all_built_in_features_are_available_without_license_flags(self):
        self.login()

        catalog = self.client.get("/api/features")
        self.assertEqual(catalog.status_code, 200)
        self.assertEqual(
            set(catalog.get_json()),
            {"features"},
        )
        self.assertIn("maintenance_schedule", catalog.get_json()["features"])
        self.assertIn("csv_export", catalog.get_json()["features"])

        created = self.client.post(
            "/api/maintenance",
            json={"device_id": self.device_id, "title": "Sicherheitsprüfung", "due_date": "2026-08-01"},
        )
        self.assertEqual(created.status_code, 201)
        summary = self.client.get("/api/maintenance/summary")
        self.assertEqual(summary.status_code, 200)
        self.assertEqual(summary.get_json()["open"], 1)
        self.assertEqual(set(summary.get_json()), {"open", "overdue"})

        export = self.client.get("/api/export/devices")
        self.assertEqual(export.status_code, 200)
        self.assertEqual(export.mimetype, "text/csv")
        self.assertIn("Testgerät", export.get_data(as_text=True))

    def test_login_lockout(self):
        for _ in range(5):
            self.client.post("/login", data={"username": "tester", "password": "wrong"})
        response = self.client.post("/login", data={"username": "tester", "password": "secret1234"})
        self.assertIn("gesperrt", response.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
