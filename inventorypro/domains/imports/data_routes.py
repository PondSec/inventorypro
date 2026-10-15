"""Protected data migration routes for the import domain."""

from __future__ import annotations

from collections.abc import Callable, Sequence
import csv
import json
from io import StringIO
from pathlib import Path
import shutil
import sqlite3
from typing import Any, Mapping
import zipfile

from flask import Blueprint, jsonify, request, session

from inventorypro.data_migration import TabularImportError
from inventorypro.web.blueprints import stable_route

from .service import ImportProfileValidationError
from .storage import stage_archive_uploads,install_staged_uploads


def build_data_import_blueprint(
    *,
    get_db: Callable[[], Any],
    get_import_settings: Callable[[Any], Mapping[str, Any]],
    should_rate_limit: Callable[[str], bool],
    current_actor: Callable[[], str],
    max_import_bytes: int,
    max_import_expanded_bytes: int,
    uploads_dir: Path,
    import_tables: Sequence[str],
    run_backup_job: Callable[..., Any],
    save_import_file: Callable[..., tuple[Path | None, str | None]],
    validate_import_file: Callable[[Path], str | None],
    resolve_tabular_import_options: Callable[..., Mapping[str, Any]],
    import_profile_service: Any,
    preview_proof_service: Any,
    build_preview_proof_arguments: Callable[..., Mapping[str, Any]],
    import_tabular_file: Callable[..., Mapping[str, Any]],
    import_data_payload: Callable[..., None],
    validate_import_archive: Callable[..., list[tuple[Any, Path]]],
    import_table_rows: Callable[..., None],
    import_from_sqlite: Callable[..., None],
    parse_tabular_file: Callable[..., Mapping[str, Any]],
    preview_tabular_file: Callable[..., dict[str, Any]],
    inspect_tabular_conflicts: Callable[..., Mapping[str, Any]],
    log_activity: Callable[..., None],
    login_required: Callable[[Callable[..., Any]], Callable[..., Any]],
    require_permission: Callable[[str], Callable[[Callable[..., Any]], Callable[..., Any]]],
) -> Blueprint:
    """Create data-import endpoints while keeping Flask contracts stable."""
    blueprint = Blueprint("data_imports", __name__)
    table_names = tuple(import_tables)

    def imports_enabled(connection: Any) -> tuple[Mapping[str, Any] | None, Any]:
        settings = get_import_settings(connection)
        if not settings["importExport"]["importAllowed"]:
            return None, (jsonify({"error": "Import ist deaktiviert."}), 403)
        return settings, None

    @stable_route(blueprint, "/api/import", methods=["POST"])
    @login_required
    @require_permission("server_settings.manage")
    def import_data():
        database = get_db()
        settings, disabled = imports_enabled(database)
        if disabled:
            return disabled
        if should_rate_limit(f"import:{current_actor()}"):
            return jsonify({"error": "Zu viele Import-Anfragen."}), 429
        import_mode = settings["importExport"]["importMode"]
        file_storage = request.files.get("file")
        file_path, error = save_import_file(
            file_storage,
            content_length=request.content_length,
            max_import_bytes=max_import_bytes,
        )
        if error:
            return jsonify({"error": error}), 400
        antivirus_error = validate_import_file(file_path)
        if antivirus_error:
            shutil.rmtree(file_path.parent, ignore_errors=True)
            return jsonify({"error": antivirus_error}), 400
        summary = None
        tabular_mode = None
        created_uploads=[]
        try:
            if import_mode == "replace":
                run_backup_job(database, settings, force=True)
            suffix = file_path.suffix.lower()
            entity = (request.form.get("entity") or "").strip().lower()
            if suffix == ".json" and entity:
                tabular_mode = (request.form.get("mode") or import_mode).strip().lower()
                content = file_path.read_bytes()
                options = resolve_tabular_import_options(database, request.form, entity, import_profile_service)
                preview_proof_service.verify(
                    request.form.get("previewToken"),
                    **build_preview_proof_arguments(content, file_path.name, entity, options, current_actor()),
                )
                summary = import_tabular_file(
                    database,
                    content,
                    entity,
                    tabular_mode,
                    file_path.name,
                    mapping_override=options["mapping"],
                    matching_key=options["matchingKey"],
                    sheet_name=options["sheetName"],
                )
            elif suffix == ".json":
                payload = json.loads(file_path.read_text(encoding="utf-8"))
                with database:
                    import_data_payload(database, payload, import_mode, table_names)
            elif suffix == ".zip":
                with zipfile.ZipFile(file_path, "r") as archive:
                    try:
                        upload_members = validate_import_archive(
                            archive,
                            max_expanded_bytes=max_import_expanded_bytes,
                        )
                    except ValueError as error:
                        return jsonify({"error": str(error)}), 400
                    members = archive.namelist()
                    data_files = [name for name in members if name.endswith('.csv') and Path(name).stem in table_names and Path(name).name==name]
                    root=uploads_dir() if callable(uploads_dir) else uploads_dir
                    staged=stage_archive_uploads(archive,upload_members,root,file_path.parent) if settings['importExport']['includeUploads'] else []
                    with database:
                        if 'inventory_export.json' in members:
                            import_data_payload(database,json.loads(archive.read('inventory_export.json')),import_mode,table_names)
                        elif 'inventory.db' in members:
                            extracted=file_path.parent/'inventory.db'
                            extracted.write_bytes(archive.read('inventory.db'))
                            import_from_sqlite(database,extracted,import_mode,table_names)
                        elif 'inventory_export.xlsx' in members:
                            from openpyxl import load_workbook
                            from io import BytesIO
                            workbook=load_workbook(BytesIO(archive.read('inventory_export.xlsx')),read_only=True,data_only=False)
                            try:
                                payload={}
                                for table in table_names:
                                    if table not in workbook.sheetnames:continue
                                    rows=workbook[table].iter_rows(values_only=True);headers=next(rows,())
                                    payload[table]=[dict(zip(headers,row)) for row in rows]
                                import_data_payload(database,payload,import_mode,table_names)
                            finally:workbook.close()
                        elif data_files:
                            payload={}
                            for name in data_files:
                                payload[Path(name).stem]=list(csv.DictReader(StringIO(archive.read(name).decode('utf-8'))))
                            import_data_payload(database,payload,import_mode,table_names)
                        else:raise ValueError('Das Archiv enthält keine unterstützten Inventardaten.')
                        created_uploads=install_staged_uploads(staged)
            elif suffix in {".csv", ".tsv", ".xlsx"}:
                tabular_mode = (request.form.get("mode") or import_mode).strip().lower()
                content = file_path.read_bytes()
                options = resolve_tabular_import_options(database, request.form, entity, import_profile_service)
                preview_proof_service.verify(
                    request.form.get("previewToken"),
                    **build_preview_proof_arguments(content, file_path.name, entity, options, current_actor()),
                )
                summary = import_tabular_file(
                    database,
                    content,
                    entity,
                    tabular_mode,
                    file_path.name,
                    mapping_override=options["mapping"],
                    matching_key=options["matchingKey"],
                    sheet_name=options["sheetName"],
                )
            elif suffix in {".db", ".sqlite"}:
                with database:
                    import_from_sqlite(database, file_path, import_mode, table_names)
            else:
                return jsonify({"error": "Unbekanntes Import-Format."}), 400
            log_activity(
                database,
                "import_completed",
                "server_settings",
                details={"mode": tabular_mode or import_mode, "entity": entity or None, "summary": summary},
            )
            database.commit()
            response = {"status": "success"}
            if suffix in {".csv", ".tsv", ".xlsx"} or (suffix == ".json" and entity):
                response["summary"] = summary
            return jsonify(response)
        except (TabularImportError, ImportProfileValidationError,ValueError,sqlite3.IntegrityError,zipfile.BadZipFile,OSError) as error:
            database.rollback()
            for path in created_uploads:path.unlink(missing_ok=True)
            log_activity(
                database,
                "import_rejected",
                "server_settings",
                details={"entity": (request.form.get("entity") or "").strip().lower() or None, "reason": str(error)},
            )
            database.commit()
            message='Der Import enthält ungültige oder weiterhin verwendete Verweise. Vorhandene Daten wurden beibehalten.' if isinstance(error,sqlite3.IntegrityError) else 'Die Importdateien konnten nicht sicher gespeichert werden. Vorhandene Dateien wurden beibehalten.' if isinstance(error,OSError) else str(error)
            return jsonify({"error": message}), 400
        finally:
            shutil.rmtree(file_path.parent, ignore_errors=True)

    @stable_route(blueprint, "/api/import/preview", methods=["POST"])
    @login_required
    @require_permission("server_settings.manage")
    def preview_import_data():
        database = get_db()
        _, disabled = imports_enabled(database)
        if disabled:
            return disabled
        file_storage = request.files.get("file")
        file_path, error = save_import_file(
            file_storage,
            content_length=request.content_length,
            max_import_bytes=max_import_bytes,
        )
        if error:
            return jsonify({"error": error}), 400
        try:
            entity = (request.form.get("entity") or "").strip().lower()
            content = file_path.read_bytes()
            options = resolve_tabular_import_options(database, request.form, entity, import_profile_service)
            parsed = parse_tabular_file(
                content,
                entity,
                file_path.name,
                mapping_override=options["mapping"],
                sheet_name=options["sheetName"],
            )
            preview = preview_tabular_file(
                content,
                entity,
                file_path.name,
                mapping_override=options["mapping"],
                sheet_name=options["sheetName"],
            )
            conflict_report = inspect_tabular_conflicts(database, parsed, options["matchingKey"])
            preview.update(conflict_report)
            preview["profileId"] = options["profileId"]
            verified_options = {
                **options,
                "mapping": parsed["mapping"],
                "matchingKey": conflict_report["matchingKey"],
                "sheetName": parsed.get("sheetName"),
            }
            preview["previewToken"] = preview_proof_service.issue(
                **build_preview_proof_arguments(content, file_path.name, entity, verified_options, current_actor()),
            )
            log_activity(
                database,
                "import_previewed",
                "server_settings",
                details={
                    "entity": entity,
                    "format": preview["format"],
                    "validRows": preview["validRows"],
                    "invalidRows": preview["invalidRows"],
                    "conflictCount": preview["conflictCount"],
                },
            )
            database.commit()
            return jsonify(preview)
        except (TabularImportError, ImportProfileValidationError) as error:
            return jsonify({"error": str(error)}), 400
        finally:
            shutil.rmtree(file_path.parent, ignore_errors=True)

    return blueprint
