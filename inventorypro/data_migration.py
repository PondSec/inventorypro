"""Safe, schema-aware migration of tabular device and asset inventories."""

from __future__ import annotations

import csv
from datetime import date, datetime
import io
import json
from pathlib import Path
import re
import sqlite3
from typing import Any, Iterable, Mapping, Sequence
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class TabularImportError(ValueError):
    """Raised for a recoverable, user-facing tabular import problem."""


MAX_TABULAR_ROWS = 100_000
MAX_XLSX_ARCHIVE_MEMBERS = 1_000
MAX_XLSX_EXPANDED_BYTES = 100 * 1024 * 1024

_ALIASES = {
    "devices": {
        "name": {"name", "device", "device name", "hostname", "host", "gerät", "geraet"},
        "serial_number": {"serial", "serial number", "serialnumber", "seriennummer", "sn"},
        "category": {"category", "device category", "kategorie", "gerätekategorie", "geraetekategorie"},
        "location": {"location", "site", "standort", "office"},
    },
    "assets": {
        "name": {"name", "asset", "asset name", "bezeichnung", "gerät", "geraet"},
        "category": {"category", "asset category", "kategorie", "assetkategorie"},
        "notes": {"notes", "note", "beschreibung", "description", "notizen"},
        "acquisition_date": {"acquisition date", "purchase date", "kaufdatum", "anschaffungsdatum"},
        "commissioning_date": {"commissioning date", "inbetriebnahme", "inbetriebnahmedatum"},
        "warranty_end": {"warranty end", "warranty expiry", "garantieende", "garantie bis"},
        "purchase_cost": {"purchase cost", "cost", "price", "preis", "anschaffungskosten"},
        "currency": {"currency", "währung", "waehrung"},
        "cost_center": {"cost center", "kostenstelle"},
        "invoice_number": {"invoice", "invoice number", "rechnungsnummer"},
    },
}

_MATCHING_KEYS = {
    "devices": {"serial_number", "name"},
    "assets": {"invoice_number", "name"},
}

_SUPPORTED_CONFLICT_STRATEGIES = {"append", "merge", "abort"}


def _key(value: str) -> str:
    normalized = value.casefold()
    normalized = normalized.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue")
    return re.sub(r"[^a-z0-9]+", " ", normalized).strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time().isoformat() == "00:00:00":
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def _validate_entity(entity: str) -> None:
    if entity not in _ALIASES:
        raise TabularImportError("Nur Geräte und Assets können tabellarisch importiert werden.")


def _build_mapping(
    entity: str,
    headers: Sequence[str],
    mapping_override: Mapping[str, str | None] | None = None,
) -> dict[str, str | None]:
    _validate_entity(entity)
    normalized_headers = [_key(header) for header in headers]
    if not any(normalized_headers):
        raise TabularImportError("Die Datei enthält keine Kopfzeile.")
    if len([header for header in normalized_headers if header]) != len(set(header for header in normalized_headers if header)):
        raise TabularImportError("Die Datei enthält doppelte Spaltenüberschriften.")
    mapping: dict[str, str | None] = {}
    for target, aliases in _ALIASES[entity].items():
        normalized_aliases = {_key(alias) for alias in aliases}
        mapping[target] = next(
            (header for header in headers if _key(header) in normalized_aliases),
            None,
        )
    if mapping_override is not None:
        if not isinstance(mapping_override, Mapping):
            raise TabularImportError("Die Feldzuordnung muss ein Objekt sein.")
        unknown_targets = set(mapping_override) - set(_ALIASES[entity])
        if unknown_targets:
            raise TabularImportError("Die Feldzuordnung enthält unbekannte Zielfelder.")
        available_headers = {header for header in headers if header}
        for target, source in mapping_override.items():
            if source is not None and source not in available_headers:
                raise TabularImportError(
                    f"Die Quellspalte für '{target}' ist in der Datei nicht vorhanden."
                )
            mapping[target] = source
    mapped_sources = [source for source in mapping.values() if source]
    if len(mapped_sources) != len(set(mapped_sources)):
        raise TabularImportError("Eine Quellspalte darf nur einem Zielfeld zugeordnet werden.")
    if not mapping["name"]:
        raise TabularImportError("Keine Spalte für Name/Gerät/Asset erkannt.")
    return mapping


def _parse_rows(
    entity: str,
    headers: Sequence[Any],
    rows: Iterable[tuple[int, Sequence[Any]]],
    *,
    source_format: str,
    delimiter: str | None = None,
    mapping_override: Mapping[str, str | None] | None = None,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    normalized_headers = [_cell_text(header) for header in headers]
    mapping = _build_mapping(entity, normalized_headers, mapping_override)
    parsed_rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for count, (line_number, values) in enumerate(rows, start=1):
        if count > MAX_TABULAR_ROWS:
            raise TabularImportError(f"Die Importdatei darf höchstens {MAX_TABULAR_ROWS:,} Datenzeilen enthalten.")
        source = {
            header: _cell_text(values[index]) if index < len(values) else ""
            for index, header in enumerate(normalized_headers)
            if header
        }
        normalized = {
            target: (source.get(header) or "").strip()
            for target, header in mapping.items()
            if header
        }
        if not normalized.get("name"):
            errors.append({"line": line_number, "error": "Name ist erforderlich."})
            continue
        normalized["specs"] = {
            header: value
            for header, value in source.items()
            if header not in mapping.values() and value
        }
        parsed_rows.append(normalized)
    return {
        "entity": entity,
        "mapping": mapping,
        "rows": parsed_rows,
        "errors": errors,
        "format": source_format,
        "delimiter": delimiter,
        "headers": normalized_headers,
        "sheetName": sheet_name,
    }


def parse_tabular_csv(
    content: bytes,
    entity: str,
    mapping_override: Mapping[str, str | None] | None = None,
) -> dict[str, Any]:
    """Parse a CSV/TSV export and infer a conservative Inventory Pro mapping."""
    _validate_entity(entity)
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise TabularImportError("Die Importdatei muss UTF-8-kodiert sein.") from error
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect=dialect)
    try:
        headers = next(reader)
    except StopIteration as error:
        raise TabularImportError("Die Datei enthält keine Kopfzeile.") from error
    return _parse_rows(
        entity,
        headers,
        enumerate(reader, start=2),
        source_format="tsv" if dialect.delimiter == "\t" else "csv",
        delimiter=dialect.delimiter,
        mapping_override=mapping_override,
    )


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ARCHIVE_MEMBERS:
                raise TabularImportError("Die XLSX-Datei enthält zu viele Archivbestandteile.")
            expanded_bytes = sum(max(0, info.file_size) for info in infos)
    except zipfile.BadZipFile as error:
        raise TabularImportError("Die XLSX-Datei ist beschädigt oder kein gültiges Archiv.") from error
    if expanded_bytes > MAX_XLSX_EXPANDED_BYTES:
        raise TabularImportError("Die entpackte XLSX-Datei überschreitet die zulässige Größe.")


def parse_tabular_xlsx(
    content: bytes,
    entity: str,
    mapping_override: Mapping[str, str | None] | None = None,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Parse a selected worksheet of a regular XLSX migration export."""
    _validate_entity(entity)
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as error:
        raise TabularImportError("Die XLSX-Datei kann nicht gelesen werden.") from error
    try:
        selected_sheet = sheet_name or (entity if entity in workbook.sheetnames else workbook.active.title)
        if selected_sheet not in workbook.sheetnames:
            raise TabularImportError("Das ausgewählte Tabellenblatt ist nicht vorhanden.")
        worksheet = workbook[selected_sheet]
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration as error:
            raise TabularImportError("Die XLSX-Datei enthält keine Kopfzeile.") from error
        parsed = _parse_rows(
            entity,
            headers,
            enumerate(rows, start=2),
            source_format="xlsx",
            mapping_override=mapping_override,
            sheet_name=selected_sheet,
        )
        parsed["availableSheets"] = list(workbook.sheetnames)
        return parsed
    finally:
        workbook.close()


def parse_tabular_json(
    content: bytes,
    entity: str,
    mapping_override: Mapping[str, str | None] | None = None,
) -> dict[str, Any]:
    """Parse a JSON array or a named devices/assets array for migration."""
    _validate_entity(entity)
    try:
        payload = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise TabularImportError("Die JSON-Datei muss UTF-8-kodiert und gültig sein.") from error
    records = payload.get(entity) if isinstance(payload, dict) else payload
    if not isinstance(records, list):
        raise TabularImportError(f"Die JSON-Datei muss ein Array oder ein '{entity}'-Array enthalten.")
    if len(records) > MAX_TABULAR_ROWS:
        raise TabularImportError(f"Die Importdatei darf höchstens {MAX_TABULAR_ROWS:,} Datenzeilen enthalten.")
    if any(not isinstance(record, Mapping) for record in records):
        raise TabularImportError("Jeder JSON-Eintrag muss ein Objekt sein.")
    headers: list[str] = []
    for record in records:
        for key in record:
            key_text = _cell_text(key)
            if key_text and key_text not in headers:
                headers.append(key_text)
    if not headers:
        raise TabularImportError("Die JSON-Datei enthält keine Kopfzeile.")
    values = [
        (index, [record.get(header) for header in headers])
        for index, record in enumerate(records, start=1)
    ]
    return _parse_rows(entity, headers, values, source_format="json", mapping_override=mapping_override)


def parse_tabular_file(
    content: bytes,
    entity: str,
    filename: str,
    *,
    mapping_override: Mapping[str, str | None] | None = None,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Parse a supported migration format based on its safely handled extension."""
    suffix = Path(filename).suffix.casefold()
    if suffix in {".csv", ".tsv"}:
        return parse_tabular_csv(content, entity, mapping_override)
    if suffix == ".xlsx":
        return parse_tabular_xlsx(content, entity, mapping_override, sheet_name)
    if suffix == ".json":
        return parse_tabular_json(content, entity, mapping_override)
    raise TabularImportError("Die Vorschau unterstützt CSV, TSV, XLSX und JSON.")


def _preview(parsed: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "entity": parsed["entity"],
        "format": parsed["format"],
        "mapping": parsed["mapping"],
        "validRows": len(parsed["rows"]),
        "invalidRows": len(parsed["errors"]),
        "errors": parsed["errors"][:100],
        "sample": parsed["rows"][:20],
        "headers": parsed["headers"],
        "sheetName": parsed.get("sheetName"),
        "availableSheets": parsed.get("availableSheets", []),
        "matchingKeys": sorted(_MATCHING_KEYS[parsed["entity"]]),
        "conflictStrategies": sorted(_SUPPORTED_CONFLICT_STRATEGIES),
    }


def preview_tabular_csv(content: bytes, entity: str) -> dict[str, Any]:
    """Return the legacy CSV/TSV preview response shape."""
    return _preview(parse_tabular_csv(content, entity))


def preview_tabular_file(
    content: bytes,
    entity: str,
    filename: str,
    *,
    mapping_override: Mapping[str, str | None] | None = None,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Return a safe preview for any supported tabular migration file."""
    return _preview(
        parse_tabular_file(
            content,
            entity,
            filename,
            mapping_override=mapping_override,
            sheet_name=sheet_name,
        )
    )


def _matching_key(parsed: Mapping[str, Any], matching_key: str | None) -> str:
    entity = parsed["entity"]
    default = "serial_number" if entity == "devices" else "invoice_number"
    selected = matching_key or (default if parsed["mapping"].get(default) else "name")
    if selected not in _MATCHING_KEYS[entity]:
        raise TabularImportError("Der Abgleichschlüssel ist für diesen Datentyp nicht zulässig.")
    if not parsed["mapping"].get(selected):
        raise TabularImportError("Der Abgleichschlüssel muss einer Quellspalte zugeordnet sein.")
    return selected


def _find_existing(connection: sqlite3.Connection, entity: str, matching_key: str, value: str | None):
    if not value:
        return None
    table = "devices" if entity == "devices" else "assets"
    return connection.execute(
        f"SELECT id FROM {table} WHERE {matching_key} = ?",
        (value,),
    ).fetchone()


def inspect_tabular_conflicts(
    connection: sqlite3.Connection,
    parsed: Mapping[str, Any],
    matching_key: str | None = None,
) -> dict[str, Any]:
    """Return bounded, row-level duplicate information without changing data."""
    selected_matching_key = _matching_key(parsed, matching_key)
    conflicts = []
    for row_number, row in enumerate(parsed["rows"], start=2):
        existing = _find_existing(
            connection,
            parsed["entity"],
            selected_matching_key,
            row.get(selected_matching_key),
        )
        if existing:
            conflicts.append({"line": row_number, "id": existing["id"], "key": selected_matching_key})
    return {
        "matchingKey": selected_matching_key,
        "conflictCount": len(conflicts),
        "conflicts": conflicts[:100],
    }


def _import_parsed(
    connection: sqlite3.Connection,
    parsed: Mapping[str, Any],
    mode: str,
    matching_key: str | None = None,
) -> dict[str, Any]:
    if mode not in _SUPPORTED_CONFLICT_STRATEGIES:
        raise TabularImportError("Importmodus muss append, merge oder abort sein.")
    if parsed["errors"]:
        raise TabularImportError("Die Vorschau enthält ungültige Zeilen; bitte zuerst korrigieren.")
    conflict_report = inspect_tabular_conflicts(connection, parsed, matching_key)
    selected_matching_key = conflict_report["matchingKey"]
    if mode == "abort" and conflict_report["conflictCount"]:
        raise TabularImportError("Der Import wurde wegen vorhandener Datensätze abgebrochen.")
    created = updated = skipped = 0
    with connection:
        for row in parsed["rows"]:
            if parsed["entity"] == "devices":
                category_id = _category_id(connection, "categories", row.get("category") or "Importiert")
                location_id = _location_id(connection, row.get("location"))
                existing = _find_existing(connection, "devices", selected_matching_key, row.get(selected_matching_key))
                if existing and mode == "append":
                    skipped += 1
                    continue
                values = (
                    row["name"], category_id, row.get("serial_number") or None,
                    location_id, json.dumps(row["specs"], ensure_ascii=False),
                )
                if existing:
                    connection.execute(
                        "UPDATE devices SET name=?, category_id=?, serial_number=?, location_id=?, specs=? WHERE id=?",
                        (*values, existing["id"]),
                    )
                    updated += 1
                else:
                    connection.execute(
                        "INSERT INTO devices (name, category_id, serial_number, location_id, specs) VALUES (?, ?, ?, ?, ?)",
                        values,
                    )
                    created += 1
            else:
                category_id = _category_id(connection, "asset_categories", row.get("category") or "Importiert")
                existing = _find_existing(connection, "assets", selected_matching_key, row.get(selected_matching_key))
                if existing and mode == "append":
                    skipped += 1
                    continue
                values = (
                    row["name"], category_id, row.get("notes") or None,
                    json.dumps(row["specs"], ensure_ascii=False), row.get("acquisition_date") or None,
                    row.get("commissioning_date") or None, row.get("warranty_end") or None,
                    _cost(row.get("purchase_cost")), row.get("currency") or None,
                    row.get("cost_center") or None, row.get("invoice_number") or None,
                )
                if existing:
                    connection.execute(
                        "UPDATE assets SET name=?, category_id=?, notes=?, specs=?, acquisition_date=?, commissioning_date=?, warranty_end=?, purchase_cost=?, currency=?, cost_center=?, invoice_number=? WHERE id=?",
                        (*values, existing["id"]),
                    )
                    updated += 1
                else:
                    connection.execute(
                        "INSERT INTO assets (name, category_id, notes, specs, acquisition_date, commissioning_date, warranty_end, purchase_cost, currency, cost_center, invoice_number) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        values,
                    )
                    created += 1
    return {"created": created, "updated": updated, "skipped": skipped, "errors": parsed["errors"]}


def import_tabular_csv(connection: sqlite3.Connection, content: bytes, entity: str, mode: str) -> dict[str, Any]:
    """Import a CSV/TSV export while preserving the established public API."""
    return _import_parsed(connection, parse_tabular_csv(content, entity), mode)


def import_tabular_file(
    connection: sqlite3.Connection,
    content: bytes,
    entity: str,
    mode: str,
    filename: str,
    *,
    mapping_override: Mapping[str, str | None] | None = None,
    matching_key: str | None = None,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Import a CSV, TSV, XLSX, or JSON migration file atomically."""
    return _import_parsed(
        connection,
        parse_tabular_file(
            content,
            entity,
            filename,
            mapping_override=mapping_override,
            sheet_name=sheet_name,
        ),
        mode,
        matching_key,
    )


def _category_id(connection: sqlite3.Connection, table: str, name: str) -> int:
    row = connection.execute(f"SELECT id FROM {table} WHERE name = ?", (name,)).fetchone()
    if row:
        return row["id"]
    return connection.execute(f"INSERT INTO {table} (name) VALUES (?)", (name,)).lastrowid


def _location_id(connection: sqlite3.Connection, name: str | None) -> int | None:
    if not name:
        return None
    row = connection.execute("SELECT id FROM locations WHERE name = ?", (name,)).fetchone()
    if row:
        return row["id"]
    return connection.execute("INSERT INTO locations (name) VALUES (?)", (name,)).lastrowid


def _cost(value: str | None) -> float | None:
    if not value:
        return None
    normalized = value.replace(" ", "")
    if "," in normalized and "." in normalized:
        decimal_separator = "," if normalized.rfind(",") > normalized.rfind(".") else "."
        thousands_separator = "." if decimal_separator == "," else ","
        normalized = normalized.replace(thousands_separator, "").replace(decimal_separator, ".")
    elif "," in normalized:
        normalized = normalized.replace(".", "").replace(",", ".")
    try:
        return float(normalized)
    except ValueError as error:
        raise TabularImportError("Anschaffungskosten sind keine gültige Zahl.") from error
